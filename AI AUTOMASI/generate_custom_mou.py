#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
GENERATOR SURAT PERJANJIAN KERJASAMA (MOU) FORMAT RESMI 1 HALAMAN
MENGIKUTI CONTOH FORMAT REFERENSI PENGGUNA
PT KEDIRI CHEMICAL ABADI — STANDAR MUTU ISO 9001:2015

Penanggung Jawab / Author : Yerikho Arfensias Effendi
Perusahaan                : PT Kediri Chemical Abadi
Direktur Utama            : Yan Effendi
"""

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

OUTPUT_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "KCA DOKUMEN", "09_DOKUMEN_DIREKSI_DAN_MANAJEMEN_PUNCAK"))
os.makedirs(OUTPUT_DIR, exist_ok=True)

COLOR_BLACK = RGBColor(0, 0, 0)

def sanitize_metadata(doc):
    core_props = doc.core_properties
    core_props.author = "Yerikho Arfensias Effendi"
    core_props.last_modified_by = "Yerikho Arfensias Effendi"
    core_props.comments = "Surat Perjanjian Kerjasama Kemitraan Kerja PT Kediri Chemical Abadi"
    core_props.category = "Surat Perjanjian Kerja"
    core_props.keywords = "PT Kediri Chemical Abadi, Surat Perjanjian Kerjasama, MOU, Farhan, Fajar, Yerikho Arfensias Effendi"

def setup_page_layout(doc):
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

def add_par(doc, text="", bold=False, italic=False, underline=False, size=11, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=6):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    if text:
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.underline = underline
        run.font.color.rgb = COLOR_BLACK
    return p

def add_identitas_table(doc, data_dict):
    tbl = doc.add_table(rows=len(data_dict), cols=3)
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl.autofit = False
    
    widths = [Inches(1.8), Inches(0.2), Inches(4.5)]
    for row_idx, (label, val) in enumerate(data_dict.items()):
        row = tbl.rows[row_idx]
        for c_idx, text in enumerate([label, ":", val]):
            cell = row.cells[c_idx]
            cell.width = widths[c_idx]
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(11)
            run.font.color.rgb = COLOR_BLACK

# -------------------------------------------------------------
# DOKUMEN 1: SURAT PERJANJIAN KERJASAMA — FARHAN
# -------------------------------------------------------------
def build_spk_farhan():
    doc = Document()
    setup_page_layout(doc)
    sanitize_metadata(doc)

    # Judul
    add_par(doc, "SURAT PERJANJIAN KERJASAMA", bold=True, underline=True, size=13, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14)
    add_par(doc, "Saya yang bertanda tangan di bawah ini :", size=11, space_after=6)

    # Pihak Pertama
    pihak_1 = {
        "Nama": "Yerikho Arfensias Effendi",
        "Jabatan": "General Manager PT Kediri Chemical Abadi",
        "Alamat": "Jl. Merbabu No. 12, Mojoroto, Kota Kediri, Jawa Timur",
        "No. KTP/SIM": "3571 . . . . . . . . . . . . . . . ."
    }
    add_identitas_table(doc, pihak_1)
    add_par(doc, "Yang mana selanjutnya akan disebut sebagai Pihak Pertama.", size=11, space_before=4, space_after=8)

    # Pihak Kedua
    pihak_2 = {
        "Nama": "Farhan",
        "Tempat, Tanggal Lahir": "Kediri, . . . . . . . . . . . . . . . . . . . .",
        "Alamat": ". . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . .",
        "No. KTP/SIM": ". . . . . . . . . . . . . . . . . . . ."
    }
    add_identitas_table(doc, pihak_2)
    add_par(doc, "Selanjutnya akan disebut dengan Pihak Kedua.", size=11, space_before=4, space_after=8)

    add_par(doc, "Kedua belah telah sepakat untuk mengadakan kerjasama usaha dengan ketentuan-ketentuan yang diatur sebagai berikut ini :", size=11, space_after=10)

    # Pasal-Pasal
    pasal_data = [
        ("PASAL 1",
         "Dalam usaha ini Pihak Pertama menunjuk Pihak Kedua untuk menjalankan tugas sebagai Business Operations Coordinator yang bertugas membantu dan mengeksekusi segala hal yang disiapkan oleh Pihak Pertama serta membantu segala kebutuhan operasional usaha PT Kediri Chemical Abadi."),

        ("PASAL 2",
         "Pihak Pertama akan memberikan salary/gaji kepada Pihak Kedua setelah arus kas perusahaan mulai stabil dan sanggup untuk membayar, dimulai secara bertahap dari Rp 1.500.000,00 (satu juta lima ratus ribu rupiah) per bulan dan akan terus dinaikkan sesuai dengan peningkatan kestabilan arus kas perusahaan."),

        ("PASAL 3",
         "Jika penjualan produk usaha telah masuk ke skala besar (produksi massal), maka sistem salary Pihak Kedua akan ditambahkan dengan pembagian komisi keuntungan sebesar 1% (satu persen) dari setiap produk yang terjual."),

        ("PASAL 4",
         "Kedua belah pihak akan saling bekerjasama untuk memajukan usaha, menjaga ketersediaan barang, serta menjaga nama baik dan kerahasiaan usaha Pihak Pertama."),

        ("PASAL 5",
         "Apabila terjadi perselisihan antar kedua belah pihak akan diselesaikan secara kekeluargaan terlebih dahulu. Dan apabila tidak ditemui jalan keluar baru akan diselesaikan secara hukum.")
    ]

    for p_title, p_desc in pasal_data:
        add_par(doc, p_title, bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=6, space_after=2)
        add_par(doc, p_desc, size=11, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=8)

    # Penutup
    add_par(doc, "Demikian surat perjanjian ini kami buat sebenar-benarnya dalam rangkap dua yang mana masing-masing rangkap mempunyai kekuatan hukum yang sama. Dan dalam pembuatan perjanjian kerjasama ini tidak ada paksaan dari pihak manapun.", size=11, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=6, space_after=12)

    # Tanggal
    add_par(doc, "Kediri, 6 September 2026", size=11, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=8)

    # Tanda Tangan
    tbl_ttd = doc.add_table(rows=1, cols=3)
    tbl_ttd.alignment = WD_TABLE_ALIGNMENT.CENTER
    c1, c2, c3 = tbl_ttd.cell(0, 0), tbl_ttd.cell(0, 1), tbl_ttd.cell(0, 2)
    c1.width = Inches(2.2)
    c2.width = Inches(2.0)
    c3.width = Inches(2.2)

    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.add_run("Pihak Pertama,\n\n\n\n\n").font.color.rgb = COLOR_BLACK
    p1.add_run("Yerikho Arfensias Effendi")

    p2 = c2.paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.add_run("\n\n(Materai 10000)\n\n").font.color.rgb = COLOR_BLACK

    p3 = c3.paragraphs[0]
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.add_run("Pihak Kedua,\n\n\n\n\n").font.color.rgb = COLOR_BLACK
    p3.add_run("Farhan")

    filename = os.path.join(OUTPUT_DIR, "SURAT_PERJANJIAN_KERJASAMA_FARHAN.docx")
    doc.save(filename)
    print(f"[OK] Saved: {filename}")

# -------------------------------------------------------------
# DOKUMEN 2: SURAT PERJANJIAN KERJASAMA — FAJAR
# -------------------------------------------------------------
def build_spk_fajar():
    doc = Document()
    setup_page_layout(doc)
    sanitize_metadata(doc)

    # Judul
    add_par(doc, "SURAT PERJANJIAN KERJASAMA", bold=True, underline=True, size=13, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14)
    add_par(doc, "Saya yang bertanda tangan di bawah ini :", size=11, space_after=6)

    # Pihak Pertama
    pihak_1 = {
        "Nama": "Yerikho Arfensias Effendi",
        "Jabatan": "General Manager PT Kediri Chemical Abadi",
        "Alamat": "Jl. Merbabu No. 12, Mojoroto, Kota Kediri, Jawa Timur",
        "No. KTP/SIM": "3571 . . . . . . . . . . . . . . . ."
    }
    add_identitas_table(doc, pihak_1)
    add_par(doc, "Yang mana selanjutnya akan disebut sebagai Pihak Pertama.", size=11, space_before=4, space_after=8)

    # Pihak Kedua
    pihak_2 = {
        "Nama": "Fajar",
        "Tempat, Tanggal Lahir": "Kediri, . . . . . . . . . . . . . . . . . . . .",
        "Alamat": ". . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . .",
        "No. KTP/SIM": ". . . . . . . . . . . . . . . . . . . ."
    }
    add_identitas_table(doc, pihak_2)
    add_par(doc, "Selanjutnya akan disebut dengan Pihak Kedua.", size=11, space_before=4, space_after=8)

    add_par(doc, "Kedua belah telah sepakat untuk mengadakan kerjasama usaha dengan ketentuan-ketentuan yang diatur sebagai berikut ini :", size=11, space_after=10)

    # Pasal-Pasal
    pasal_data = [
        ("PASAL 1",
         "Dalam usaha ini Pihak Pertama menunjuk Pihak Kedua sebagai Head of Creative & Design untuk mengerjakan seluruh kebutuhan identitas visual, desain kemasan produk baru, banner website profesional, feed media sosial, dan materi promosi PT Kediri Chemical Abadi."),

        ("PASAL 2",
         "Pihak Kedua akan bekerja dengan sistem volunteer di awal proyek sampai arus kas perusahaan sanggup untuk membayar. Setiap pekerjaan dan proyek yang diselesaikan oleh Pihak Kedua akan dicatat dan disimpan buktinya oleh kedua belah pihak dalam Buku Log Proyek (Project Ledger)."),

        ("PASAL 3",
         "Penilaian harga jasa desain bersifat fleksibel tergantung dari peminat di pasar dan dampak (impact) yang dibawa oleh desain terhadap penjualan produk. Setelah arus kas perusahaan sanggup membayar, Pihak Pertama akan membayar proyek-proyek yang telah dikerjakan sebelumnya secara bertahap satu per satu disesuaikan dengan arus kas dan dana alokasi perusahaan."),

        ("PASAL 4",
         "Kedua belah pihak akan saling bekerjasama untuk mempromosikan hasil produk, dan seluruh aset desain final yang telah diserahterimakan menjadi hak guna komersial usaha Pihak Pertama, dengan Pihak Kedua tetap memiliki hak mencantumkannya dalam portofolio pribadi."),

        ("PASAL 5",
         "Apabila terjadi perselisihan antar kedua belah pihak akan diselesaikan secara kekeluargaan terlebih dahulu. Dan apabila tidak ditemui jalan keluar baru akan diselesaikan secara hukum.")
    ]

    for p_title, p_desc in pasal_data:
        add_par(doc, p_title, bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=6, space_after=2)
        add_par(doc, p_desc, size=11, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=8)

    # Penutup
    add_par(doc, "Demikian surat perjanjian ini kami buat sebenar-benarnya dalam rangkap dua yang mana masing-masing rangkap mempunyai kekuatan hukum yang sama. Dan dalam pembuatan perjanjian kerjasama ini tidak ada paksaan dari pihak manapun.", size=11, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=6, space_after=12)

    # Tanggal
    add_par(doc, "Kediri, 6 September 2026", size=11, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=8)

    # Tanda Tangan
    tbl_ttd = doc.add_table(rows=1, cols=3)
    tbl_ttd.alignment = WD_TABLE_ALIGNMENT.CENTER
    c1, c2, c3 = tbl_ttd.cell(0, 0), tbl_ttd.cell(0, 1), tbl_ttd.cell(0, 2)
    c1.width = Inches(2.2)
    c2.width = Inches(2.0)
    c3.width = Inches(2.2)

    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.add_run("Pihak Pertama,\n\n\n\n\n").font.color.rgb = COLOR_BLACK
    p1.add_run("Yerikho Arfensias Effendi")

    p2 = c2.paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.add_run("\n\n(Materai 10000)\n\n").font.color.rgb = COLOR_BLACK

    p3 = c3.paragraphs[0]
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.add_run("Pihak Kedua,\n\n\n\n\n").font.color.rgb = COLOR_BLACK
    p3.add_run("Fajar")

    filename = os.path.join(OUTPUT_DIR, "SURAT_PERJANJIAN_KERJASAMA_FAJAR.docx")
    doc.save(filename)
    print(f"[OK] Saved: {filename}")

if __name__ == "__main__":
    build_spk_farhan()
    build_spk_fajar()
