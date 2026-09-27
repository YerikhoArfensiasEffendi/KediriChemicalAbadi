import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
from kca_doc_humanizer import sanitize_docx_metadata

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def generate_tax_power_of_attorney():
    base_dir = "/Users/arthur/Documents/WebsiteKCA/kediri-chemical"
    output_dir = os.path.join(base_dir, "KCA DOKUMEN/08_LEGALITAS_HR_PJT_DAN_ADMINISTRASI_MAKLOON")
    os.makedirs(output_dir, exist_ok=True)
    doc_path = os.path.join(output_dir, "SURAT_KUASA_KHUSUS_PENGURUSAN_PAJAK_KCA.docx")

    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.85)
        s.bottom_margin = Inches(0.85)
        s.left_margin = Inches(0.85)
        s.right_margin = Inches(0.85)

    # Kop Surat Resmi KCA
    kop = doc.add_paragraph()
    kop.alignment = WD_ALIGN_PARAGRAPH.CENTER
    k1 = kop.add_run("PT KEDIRI CHEMICAL ABADI\n")
    k1.font.name = "Arial"
    k1.font.size = Pt(15)
    k1.font.bold = True
    k1.font.color.rgb = RGBColor(0, 0, 0)

    k2 = kop.add_run("Pabrik: RT.1/RW.6, Pagung, Kec. Semen, Kabupaten Kediri, Jawa Timur 64161\n"
                    "Hotline/WhatsApp: 0822-4400-6699 | Email: kdrchemicals@gmail.com\n")
    k2.font.name = "Arial"
    k2.font.size = Pt(9)
    k2.font.color.rgb = RGBColor(0, 0, 0)

    p_line = doc.add_paragraph("═" * 58)
    p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in p_line.runs:
        r.font.color.rgb = RGBColor(0, 0, 0)

    # Judul Dokumen
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t1 = p_title.add_run("SURAT KUASA KHUSUS PERPAJAKAN\n")
    t1.font.name = "Arial"
    t1.font.size = Pt(13)
    t1.font.bold = True
    t1.font.color.rgb = RGBColor(0, 0, 0)

    t2 = p_title.add_run("Nomor: 016/SK-DIR/TAX/IX/2026\n(Sesuai Ketentuan Pasal 32 UU KUP & PMK No. 229/PMK.03/2014)\n\n")
    t2.font.name = "Arial"
    t2.font.size = Pt(9.5)
    t2.font.italic = True
    t2.font.color.rgb = RGBColor(0, 0, 0)

    # Pembuka
    p_buka = doc.add_paragraph("Yang bertanda tangan di bawah ini:")
    p_buka.paragraph_format.line_spacing = 1.15
    for r in p_buka.runs:
        r.font.name = "Arial"
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0, 0, 0)

    # Tabel Pemberi Kuasa
    t_pemberi = doc.add_table(rows=6, cols=2)
    t_pemberi.alignment = WD_TABLE_ALIGNMENT.CENTER
    pemberi_data = [
        ("Nama Lengkap", ": YAN EFFENDI"),
        ("NIK / No. KTP", ": (Diisi Nomor KTP Bpk. Yan Effendi)"),
        ("NPWP Pribadi", ": (Diisi Nomor NPWP Bpk. Yan Effendi)"),
        ("Jabatan", ": Direktur Utama"),
        ("Nama Badan Usaha", ": PT KEDIRI CHEMICAL ABADI"),
        ("NPWP Badan", ": (Diisi Nomor NPWP PT Kediri Chemical Abadi)")
    ]
    for idx, (label, val) in enumerate(pemberi_data):
        row = t_pemberi.rows[idx].cells
        row[0].text = label
        row[1].text = val
        for c in row:
            set_cell_margins(c)
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.name = "Arial"
                    r.font.size = Pt(9.5)
                    r.font.color.rgb = RGBColor(0, 0, 0)

    doc.add_paragraph()
    p_mid = doc.add_paragraph(
        "Bertindak selaku Direktur Utama dan Penanggung Jawab sah Wajib Pajak Badan PT KEDIRI CHEMICAL ABADI, "
        "dengan ini memberikan kuasa khusus kepada:"
    )
    p_mid.paragraph_format.line_spacing = 1.15
    for r in p_mid.runs:
        r.font.name = "Arial"
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0, 0, 0)

    # Tabel Penerima Kuasa
    t_penerima = doc.add_table(rows=4, cols=2)
    t_penerima.alignment = WD_TABLE_ALIGNMENT.CENTER
    penerima_data = [
        ("Nama Lengkap", ": YERIKHO ARFENSIAS EFFENDI"),
        ("NIK / No. KTP", ": (Diisi Nomor KTP Bpk. Yerikho)"),
        ("Jabatan", ": General Manager / Penanggung Jawab Keuangan & Operasional"),
        ("Alamat", ": RT.1/RW.6, Pagung, Kec. Semen, Kabupaten Kediri, Jawa Timur 64161")
    ]
    for idx, (label, val) in enumerate(penerima_data):
        row = t_penerima.rows[idx].cells
        row[0].text = label
        row[1].text = val
        for c in row:
            set_cell_margins(c)
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.name = "Arial"
                    r.font.size = Pt(9.5)
                    r.font.color.rgb = RGBColor(0, 0, 0)

    # Ruang Lingkup Khusus Kuasa
    doc.add_paragraph()
    p_khusus = doc.add_paragraph("-------------------------------------------------- KHUSUS --------------------------------------------------")
    p_khusus.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in p_khusus.runs:
        r.font.name = "Arial"
        r.font.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0, 0, 0)

    p_isi_kuasa = doc.add_paragraph(
        "Untuk dan atas nama Pemberi Kuasa serta Wajib Pajak Badan PT KEDIRI CHEMICAL ABADI menghadap pejabat berwenang "
        "di Kantor Pelayanan Pajak (KPP) Pratama Kediri guna melaksanakan urusan administrasi perpajakan sebagai berikut:\n\n"
        "1. Mengajukan dan menandatangani tanda terima penyerahan berkas Permohonan Pengurangan atau Penghapusan Sanksi Administrasi "
        "Perpajakan atas Surat Tagihan Pajak (STP) berdasarkan ketentuan Pasal 36 Ayat (1) Huruf a Undang-Undang KUP;\n\n"
        "2. Mengajukan dan menandatangani berkas Permohonan Penetapan Wajib Pajak Non-Efektif (WP NE) PT Kediri Chemical Abadi "
        "berdasarkan ketentuan PER-04/PJ/2020;\n\n"
        "3. Memberikan penjelasan, klarifikasi, dan menyerahkan dokumen kelengkapan pendukung legalitas perusahaan yang diminta "
        "oleh petugas Tempat Pelayanan Terpadu (TPT) maupun Account Representative (AR) KPP Pratama Kediri;\n\n"
        "4. Menandatangani dan menerima Bukti Penerimaan Surat (BPS), Surat Keputusan (SK), serta surat-menyurat resmi lainnya "
        "yang diterbitkan oleh Kantor Pelayanan Pajak Pratama Kediri terkait permohonan tersebut.\n\n"
        "Surat kuasa ini diberikan dengan hak substitusi terbatas dan berlaku sejak tanggal ditandatangani sampai dengan selesainya "
        "proses administrasi penetapan Wajib Pajak Non-Efektif dan penghapusan sanksi denda pajak tersebut di atas."
    )
    p_isi_kuasa.paragraph_format.line_spacing = 1.15
    for r in p_isi_kuasa.runs:
        r.font.name = "Arial"
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0, 0, 0)

    # Tanda Tangan
    doc.add_paragraph("\n")
    p_tgl = doc.add_paragraph("Kediri, 05 September 2026")
    p_tgl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in p_tgl.runs:
        r.font.name = "Arial"
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0, 0, 0)

    t_sign = doc.add_table(rows=3, cols=2)
    t_sign.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_sign.rows[0].cells[0].text = "Penerima Kuasa,\nGENERAL MANAGER"
    t_sign.rows[0].cells[1].text = "Pemberi Kuasa,\nDIREKTUR UTAMA"
    t_sign.rows[1].cells[0].text = "\n\n\n\n"
    t_sign.rows[1].cells[1].text = "\n\n(Meterai Rp 10.000 & Cap PT)\n\n"
    t_sign.rows[2].cells[0].text = "YERIKHO ARFENSIAS EFFENDI\nNIK: (No. KTP Bpk. Yerikho)"
    t_sign.rows[2].cells[1].text = "YAN EFFENDI\nNIK: (No. KTP Bpk. Yan Effendi)"

    for row in t_sign.rows:
        for c in row.cells:
            set_cell_margins(c)
            for p in c.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.name = "Arial"
                    r.font.bold = True
                    r.font.size = Pt(9.5)
                    r.font.color.rgb = RGBColor(0, 0, 0)

    doc.save(doc_path)
    
    # Salin ke root folder KCA DOKUMEN jika ada
    ext_path = "/Users/arthur/Documents/KCA DOKUMEN/08_LEGALITAS_HR_PJT_DAN_ADMINISTRASI_MAKLOON/SURAT_KUASA_KHUSUS_PENGURUSAN_PAJAK_KCA.docx"
    if os.path.exists(os.path.dirname(ext_path)):
        doc.save(ext_path)
        sanitize_docx_metadata(ext_path, title="Surat Kuasa Khusus Pengurusan Pajak KCA", author="Yerikho Arfensias Effendi")

    sanitize_docx_metadata(doc_path, title="Surat Kuasa Khusus Pengurusan Pajak KCA", author="Yerikho Arfensias Effendi")
    print("SUCCESS: Tax Power of Attorney generated at", doc_path)
    return doc_path

if __name__ == "__main__":
    generate_tax_power_of_attorney()
