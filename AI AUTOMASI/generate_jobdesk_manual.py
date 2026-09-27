#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
GENERATOR BUKU MANUAL LENGKAP URAIAN JABATAN & SOP OPERASIONAL TIM
PT KEDIRI CHEMICAL ABADI — STANDAR MUTU ISO 9001:2015
Edisi Tata Kelola: Single Reporting Bridge & Supreme Approver Gatekeeping

Penanggung Jawab / Author : Yerikho Arfensias Effendi
Perusahaan                : PT Kediri Chemical Abadi
Direktur Utama            : Yan Effendi
"""

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

OUTPUT_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "KCA DOKUMEN", "09_DOKUMEN_DIREKSI_DAN_MANAJEMEN_PUNCAK"))
os.makedirs(OUTPUT_DIR, exist_ok=True)

COLOR_BLACK = RGBColor(0, 0, 0)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}>\n'
                          f'<w:top w:val="{kwargs.get("top", "single")}" w:sz="{kwargs.get("top_sz", "4")}" w:space="0" w:color="{kwargs.get("top_color", "000000")}"/>\n'
                          f'<w:left w:val="{kwargs.get("left", "single")}" w:sz="{kwargs.get("left_sz", "4")}" w:space="0" w:color="{kwargs.get("left_color", "000000")}"/>\n'
                          f'<w:bottom w:val="{kwargs.get("bottom", "single")}" w:sz="{kwargs.get("bottom_sz", "4")}" w:space="0" w:color="{kwargs.get("bottom_color", "000000")}"/>\n'
                          f'<w:right w:val="{kwargs.get("right", "single")}" w:sz="{kwargs.get("right_sz", "4")}" w:space="0" w:color="{kwargs.get("right_color", "000000")}"/>\n'
                          f'</w:tcBorders>')
    tcPr.append(tcBorders)

def sanitize_metadata(doc):
    core_props = doc.core_properties
    core_props.author = "Yerikho Arfensias Effendi"
    core_props.last_modified_by = "Yerikho Arfensias Effendi"
    core_props.comments = "Manual Uraian Jabatan & SOP Operasional Tim PT Kediri Chemical Abadi Standar ISO 9001:2015"
    core_props.category = "Dokumen Manajemen Operasional"
    core_props.keywords = "PT Kediri Chemical Abadi, Jobdesk, SOP, Operasional, Yan Effendi, Yerikho, Farhan, Fajar"

def setup_page_layout(doc):
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

def add_header_kop(doc, doc_title, doc_number):
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    cell_left = tbl.cell(0, 0)
    cell_right = tbl.cell(0, 1)
    cell_left.width = Inches(4.5)
    cell_right.width = Inches(2.3)

    p_left = cell_left.paragraphs[0]
    p_left.paragraph_format.space_after = Pt(2)
    p_left.paragraph_format.line_spacing = 1.15
    run1 = p_left.add_run("PT KEDIRI CHEMICAL ABADI\n")
    run1.font.name = 'Times New Roman'
    run1.font.size = Pt(12)
    run1.font.bold = True
    run1.font.color.rgb = COLOR_BLACK

    run2 = p_left.add_run("Pusat Riset Formulasi & Manufaktur Kimia Pembersih Industri\nJl. Merbabu No. 12, Mojoroto, Kota Kediri, Jawa Timur | Telp: 0858-1230-7629")
    run2.font.name = 'Times New Roman'
    run2.font.size = Pt(8.5)
    run2.font.color.rgb = COLOR_BLACK

    p_right = cell_right.paragraphs[0]
    p_right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_right.paragraph_format.space_after = Pt(2)
    p_right.paragraph_format.line_spacing = 1.15
    run3 = p_right.add_run(f"STANDAR ISO 9001:2015\nNo: {doc_number}\nEdisi: September 2026")
    run3.font.name = 'Times New Roman'
    run3.font.size = Pt(8)
    run3.font.bold = True
    run3.font.color.rgb = COLOR_BLACK

    for c in (cell_left, cell_right):
        set_cell_border(c, top="none", left="none", right="none", bottom="single", bottom_sz="12", bottom_color="000000")
        set_cell_margins(c, top=40, bottom=80, left=40, right=40)

    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_after = Pt(6)

def add_par(doc, text="", bold=False, italic=False, size=10, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=6):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(space_after)
    if text:
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = COLOR_BLACK
    return p

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = COLOR_BLACK
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(10.5)
    run.font.bold = True
    run.font.color.rgb = COLOR_BLACK
    return p

def build_jobdesk_manual():
    doc = Document()
    setup_page_layout(doc)
    sanitize_metadata(doc)
    add_header_kop(doc, "MANUAL URAIAN JABATAN & SOP", "MGT-07/SOP-JOBDESK/2026-V2")

    add_par(doc, "BUKU PANDUAN LENGKAP URAIAN JABATAN (JOB DESCRIPTION)", bold=True, size=13, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_par(doc, "DAN STANDAR OPERASIONAL PROSEDUR (SOP) ALUR KERJA TIM INTI", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_par(doc, "PT KEDIRI CHEMICAL ABADI — TAHUN OPERASIONAL 2026", bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14)

    # BAB I
    add_heading_1(doc, "BAB I: STRUKTUR TATA KELOLA & ALUR KOMUNIKASI SATU PINTU")
    add_par(doc, "PT Kediri Chemical Abadi menerapkan sistem tata kelola eksekutif 'Two-Tier Governance & Single Reporting Bridge'. Direktur Utama / Owner bertindak sebagai pengawas tertinggi dan pemutus mutlak (Supreme Reviewer & Approver) yang hanya menerima laporan berkala dan proposal proyek dari General Manager.")
    add_par(doc, "Staf operasional (Farhan) dan divisi kreatif (Fajar) berada di bawah komando langsung General Manager (Yerikho) dan TIDAK melapor langsung kepada Direktur Utama. Hal ini untuk menjaga fokus Direktur Utama pada pengawasan strategis dan memastikan kepemimpinan operasional General Manager berjalan solid dan efektif.")

    # BAB II: JOBDESK RINCI
    add_heading_1(doc, "BAB II: URAIAN TUGAS DAN TANGGUNG JAWAB RINCI (JOB DESCRIPTION)")

    # 1. DIREKTUR UTAMA
    add_heading_2(doc, "1. DIREKTUR UTAMA / OWNER (YAN EFFENDI)")
    add_par(doc, "Kedudukan Hierarki: Pemegang Otoritas Tertinggi & Pengawas Korporat (Supreme Reviewer & Final Approver).")
    add_par(doc, "Fungsi Utama: Menerima laporan komprehensif, mengkaji kelayakan usulan proyek dari General Manager, dan memberikan keputusan mutlak APPROVE atau NOT.")
    add_par(doc, "Tugas Pokok:", bold=True, space_after=2)
    add_par(doc, "a. Eksekutif Review & Approval Proyek: Menerima laporan perkembangan operasional dan proposal proyek baru yang diajukan oleh General Manager (Yerikho). Melakukan kajian strategis, lalu memberikan keputusan mutlak: APPROVE (Disetujui untuk Dijalankan) atau NOT (Ditolak / Ditunda untuk Disempurnakan).\n"
                 "b. Pelaksanaan Proyek: Proyek baru, pengadaan mesin reaktor, atau ekspansi lini produk HANYA DAPAT BERJALAN apabila telah mendapatkan status APPROVE tertulis dari Direktur Utama.\n"
                 "c. Supervisi Teknis Formulasi & Mutu: Mengawasi integritas formulasi kimia, memberikan bimbingan teknis laboratorium, dan memastikan formula mematuhi standar mutu ISO 9001:2015 serta perizinan edar PKRT / BPOM.\n"
                 "d. Evaluasi Triwulan: Menerima rekapitulasi laba rugi triwulan dari General Manager sebagai dasar pembagian dividen modal.")
    add_par(doc, "Wewenang Mutlak: Hak veto pembatalan proyek, hak persetujuan akhir kelayakan jalan proyek (Approve / Reject), dan otorisasi kebijakan umum perusahaan.")
    add_par(doc, "Batas Wewenang: Tidak menangani komunikasi harian dengan Farhan dan Fajar, tidak mengurus packing pesanan eceran, logistik vendor, atau materi media sosial.")

    # 2. GENERAL MANAGER
    add_heading_2(doc, "2. GENERAL MANAGER (YERIKHO ARFENSIAS EFFENDI)")
    add_par(doc, "Kedudukan Hierarki: Pemimpin Tertinggi Operasional Harian & Satu-Satunya Penghubung ke Direktur Utama (Single Reporting Bridge).")
    add_par(doc, "Fungsi Utama: Merancang proposal proyek, menyusun laporan berkala kepada Direktur Utama, dan memimpin eksekusi tim Farhan dan Fajar setelah proyek berstatus APPROVED.")
    add_par(doc, "Tugas Pokok:", bold=True, space_after=2)
    add_par(doc, "a. Penyusunan Proposal Proyek ke Direktur Utama: Menyiapkan usulan bisnis terperinci (target produk, HPP, estimasi omzet, alokasi anggaran, dan jadwal kerja) lalu mempresentasikannya kepada Direktur Utama untuk mendapatkan status APPROVE atau NOT.\n"
                 "b. Komando Eksekusi Lapangan: Setelah proyek berstatus APPROVED, General Manager memegang kendali penuh membagi tugas kepada Farhan (operasional) dan Fajar (kreatif) serta memastikan target tercapai.\n"
                 "c. Pengendalian Keuangan & Arus Kas: Menghitung HPP, menetapkan harga jual, mengontrol kas masuk/keluar, dan mengalokasikan pencairan hak gaji Farhan dan hak proyek Fajar.\n"
                 "d. Validasi Approval Internal: Memeriksa dan menyetujui hasil desain kemasan dari Fajar sebelum naik cetak serta menyetujui anggaran belanja operasional dari Farhan.\n"
                 "e. Pelaporan Berkala: Menyusun dan menyerahkan laporan pertanggungjawaban operasional dan keuangan kepada Direktur Utama secara rutin.")
    add_par(doc, "Wewenang Penuh: Otorisasi seluruh operasional harian, pengeluaran kas belanja, persetujuan cetak desain, dan penugasan kerja kepada tim.")

    # 3. BUSINESS OPERATIONS COORDINATOR
    add_heading_2(doc, "3. BUSINESS OPERATIONS COORDINATOR (FARHAN)")
    add_par(doc, "Kedudukan Hierarki: Bertanggung jawab langsung 100% kepada General Manager (Yerikho).")
    add_par(doc, "Tujuan Jabatan: Tangan kanan GM yang bertindak sebagai mesin penggerak rantai pasok fisik, pengadaan bahan kemasan, pengisian botol, dan kelancaran pengiriman pesanan.")
    add_par(doc, "Tugas Pokok:", bold=True, space_after=2)
    add_par(doc, "a. Pengadaan Fisik (Procurement): Mencari sampel botol (250ml/500ml), mengukur dimensi stiker secara presisi untuk Fajar, dan mengoordinasikan percetakan stiker vinyl anti-air.\n"
                 "b. Operasional Pengisian & Stok: Membantu proses pengisian cairan sabun, memasang stiker label botol dengan rapi, dan mencatat stok barang jadi.\n"
                 "c. Marketplace & Customer Service: Membuka dan mengelola toko resmi Shopee/Tokopedia, mengunggah foto produk Fajar, dan membalas chat pembeli secara cepat.\n"
                 "d. Pengepakan & Logistik: Membungkus pesanan dengan bubble wrap aman, menempel resi pengiriman, dan menyerahkan paket ke ekspedisi (J&T/SiCepat).\n"
                 "e. Pelaporan Harian: Menyerahkan rekap pesanan terkirim dan sisa stok kepada General Manager setiap sore.")
    add_par(doc, "Garis Komunikasi: Hanya menerima instruksi dari dan melapor kepada General Manager.")

    # 4. HEAD OF CREATIVE & DESIGN
    add_heading_2(doc, "4. HEAD OF CREATIVE & DESIGN (FAJAR)")
    add_par(doc, "Kedudukan Hierarki: Bertanggung jawab langsung 100% kepada General Manager (Yerikho).")
    add_par(doc, "Tujuan Jabatan: Pemimpin estetika merek yang bertanggung jawab merancang desain kemasan mewah, materi visual website Gen-Z, dan konten media sosial yang memikat.")
    add_par(doc, "Tugas Pokok:", bold=True, space_after=2)
    add_par(doc, "a. Desain Kemasan Siap Cetak: Merancang label botol sabun mandi & shampoo sesuai ukuran presisi Farhan dengan standar percetakan (PDF Vektor/AI, CMYK, 300 DPI, Bleed potong).\n"
                 "b. Aset Visual Digital: Mendesain banner website resmi, mock-up botol 3D, dan template etalase Shopee/Tokopedia (1080x1080 px).\n"
                 "c. Konten Media Sosial: Merancang 9 feed Instagram estetik Gen-Z, template Story promosi, dan materi promosi peluncuran produk baru.\n"
                 "d. Pencatatan Project Ledger: Mendokumentasikan setiap file desain yang selesai ke formulir Project Ledger bersama General Manager untuk hak pencairan kompensasi.")
    add_par(doc, "Garis Komunikasi: Hanya menerima instruksi dari dan melapor kepada General Manager.")

    # BAB III: SOP ALUR KERJA LINTAS FUNGSI
    add_heading_1(doc, "BAB III: STANDAR OPERASIONAL PROSEDUR (SOP) GERBANG APPROVAL & EKSEKUSI")

    add_heading_2(doc, "SOP-01: ALUR GERBANG KEPUTUSAN PROYEK (GATEKEEPING APPROVAL WORKFLOW)")
    sop_steps = [
        ("Tahap 1: Inisiasi & Riset Bisnis (GM)", "Yerikho menyusun proposal proyek lini produk baru (konsep sabun mandi Gen-Z, HPP, estimasi omzet, dan alokasi modal)."),
        ("Tahap 2: Pengajuan Laporan ke Direktur Utama (GM -> Dirut)", "Yerikho mempresentasikan proposal komprehensif kepada Yan Effendi (Direktur Utama / Owner)."),
        ("Tahap 3: Executive Review & Keputusan Mutlak (Dirut)", "Yan Effendi meninjau proposal dan memberikan keputusan: APPROVE or NOT.\n"
         "  • Jika NOT  : Proyek ditunda atau direvisi sesuai catatan perbaikan Dirut.\n"
         "  • Jika APPROVE : Proyek RESMI BERJALAN (Lampu Hijau Eksekusi)."),
        ("Tahap 4: Delegasi Tugas Operasional (GM -> Tim)", "Setelah status APPROVED, Yerikho menugaskan Farhan (survei botol & percetakan) dan Fajar (desain kemasan & visual)."),
        ("Tahap 5: Supervisi Mutu Formula (Dirut & GM)", "Yan Effendi meracik formula sabun di laboratorium reaktor pabrik."),
        ("Tahap 6: Eksekusi Produksi & Kemasan (Ops & Design)", "Farhan mengukur botol -> Fajar mendesain label -> Yerikho ACC desain -> Farhan cetak stiker & isi sabun."),
        ("Tahap 7: Pemasaran & Penjualan (Ops & Design)", "Fajar memproduksi materi promosi medsos -> Farhan memasang produk di Shopee & mengurus pengiriman paket."),
        ("Tahap 8: Laporan Pertanggungjawaban Akhir (GM -> Dirut)", "Yerikho merekap hasil penjualan, laba bersih, dan menyajikan laporan berkala kepada Direktur Utama.")
    ]
    for stitle, sdesc in sop_steps:
        p = add_par(doc)
        p.add_run(f"{stitle}: ").bold = True
        p.add_run(sdesc)

    # BAB IV: MATRIKS OTORISASI KEPUTUSAN
    add_heading_1(doc, "BAB IV: MATRIKS OTORISASI KEPUTUSAN (AUTHORITY MATRIX)")
    
    headers = ["Jenis Keputusan / Aktivitas", "Pelaksana", "Reviewer / Approver", "Dampak Jika 'NOT'"]
    widths = [Inches(2.5), Inches(1.3), Inches(1.6), Inches(1.8)]
    tbl_auth = doc.add_table(rows=6, cols=4)
    tbl_auth.alignment = WD_TABLE_ALIGNMENT.CENTER

    auth_data = [
        ("Jenis Keputusan / Aktivitas", "Pelaksana", "Reviewer / Approver", "Dampak Jika 'NOT'"),
        ("Persetujuan Peluncuran Proyek Baru", "Yerikho (GM)", "Yan Effendi (Dirut)", "Proyek tidak boleh jalan"),
        ("Perubahan Formulasi Kimia Sabun", "Yan Effendi (Dirut)", "Yan Effendi (Dirut)", "Formula lama dipertahankan"),
        ("Persetujuan Desain Kemasan Naik Cetak", "Fajar (Design)", "Yerikho (GM)", "Desain wajib direvisi Fajar"),
        ("Otorisasi Belanja Botol & Stiker", "Farhan (Ops)", "Yerikho (GM)", "Dana belanja tidak cair"),
        ("Penetapan Harga Jual & Diskon Toko", "Yerikho (GM)", "Yerikho (GM)", "Harga standar berlaku")
    ]

    for r_idx, row in enumerate(auth_data):
        for c_idx, val in enumerate(row):
            cell = tbl_auth.rows[r_idx].cells[c_idx]
            cell.width = widths[c_idx]
            p = cell.paragraphs[0]
            if r_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(val)
                run.bold = True
                run.font.size = Pt(8.5)
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = p.add_run(val)
                run.font.size = Pt(8)
            run.font.name = 'Times New Roman'
            run.font.color.rgb = COLOR_BLACK
            set_cell_margins(cell, top=60, bottom=60, left=40, right=40)
            set_cell_border(cell, top="single", bottom="single", left="single", right="single", top_sz="4", bottom_sz="4")

    add_par(doc, "", space_after=14)

    # Tanda Tangan Pengesahan
    add_heading_1(doc, "BAB V: PENGESAHAN DOKUMEN MANUAL KERJA")
    add_par(doc, "Dokumen Uraian Jabatan dan SOP Gerbang Approval ini disahkan sebagai pedoman tata kelola resmi yang mengikat seluruh jajaran manajemen PT Kediri Chemical Abadi terhitung sejak tanggal 6 September 2026.", space_after=14)

    tbl_ttd = doc.add_table(rows=2, cols=2)
    tbl_ttd.alignment = WD_TABLE_ALIGNMENT.CENTER
    c1, c2 = tbl_ttd.cell(0, 0), tbl_ttd.cell(0, 1)
    c3, c4 = tbl_ttd.cell(1, 0), tbl_ttd.cell(1, 1)

    for c in (c1, c2, c3, c4):
        set_cell_border(c, top="none", left="none", right="none", bottom="none")
        c.width = Inches(3.4)

    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.add_run("Disetujui Oleh,\nDirektur Utama / Owner\n(Supreme Reviewer & Approver)\n\n\n\n").font.color.rgb = COLOR_BLACK
    p1.add_run("YAN EFFENDI").bold = True

    p2 = c2.paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.add_run("Ditetapkan di Kota Kediri, 6 September 2026\nGeneral Manager\n(Operations & Executive Lead)\n\n\n\n").font.color.rgb = COLOR_BLACK
    p2.add_run("YERIKHO ARFENSIAS EFFENDI").bold = True

    p3 = c3.paragraphs[0]
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.add_run("\n\nPihak Yang Melaksanakan,\nBusiness Operations Coordinator\n\n\n\n").font.color.rgb = COLOR_BLACK
    p3.add_run("FARHAN").bold = True

    p4 = c4.paragraphs[0]
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p4.add_run("\n\nPihak Yang Melaksanakan,\nHead of Creative & Design\n\n\n\n").font.color.rgb = COLOR_BLACK
    p4.add_run("FAJAR").bold = True

    filename = os.path.join(OUTPUT_DIR, "MGT-07_Manual_Lengkap_Jobdesk_dan_SOP_Operasional_Tim_KCA.docx")
    doc.save(filename)
    print(f"[OK] Master Jobdesk & SOP Manual updated with Approval Gate: {filename}")

if __name__ == "__main__":
    build_jobdesk_manual()
